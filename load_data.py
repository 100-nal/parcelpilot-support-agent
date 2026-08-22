"""
Loads ParcelPilot_Assessment_Data.xlsx into the Neon Postgres tables:
accounts, orders, tickets.

Usage:
    export DATABASE_URL="postgresql://neondb_owner:...@.../neondb?sslmode=require"
    python load_data.py /path/to/ParcelPilot_Assessment_Data.xlsx
"""

import os
import sys
import psycopg2
import openpyxl

DATABASE_URL = os.environ["DATABASE_URL"]

TABLE_COLUMNS = {
    "accounts": [
        "account_id", "account_name", "plan", "status",
        "csm", "contract_file", "premium_support", "notes",
    ],
    "orders": [
        "order_id", "account_id", "carrier", "status", "booked_at",
        "pickup_window_start", "pickup_window_end", "pickup_actual_at",
        "shipment_fee_inr", "carrier_fault", "customer_fault",
        "cancellation_requested_at", "notes",
    ],
    "tickets": [
        "ticket_id", "account_id", "created_at", "status", "subject",
        "description", "channel", "assigned_to",
        "last_customer_message_at", "historical_resolution",
    ],
}


def load_sheet(cur, wb, sheet_name):
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    header, data_rows = rows[0], rows[1:]

    cols = TABLE_COLUMNS[sheet_name]
    col_idx = [header.index(c) for c in cols]

    placeholders = ", ".join(["%s"] * len(cols))
    col_list = ", ".join(cols)
    sql = f"INSERT INTO {sheet_name} ({col_list}) VALUES ({placeholders}) " \
          f"ON CONFLICT DO NOTHING"

    count = 0
    for row in data_rows:
        if row[0] is None:
            continue
        values = [row[i] for i in col_idx]
        cur.execute(sql, values)
        count += 1
    print(f"Loaded {count} rows into {sheet_name}")


def main():
    xlsx_path = sys.argv[1] if len(sys.argv) > 1 else "ParcelPilot_Assessment_Data.xlsx"
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    conn = psycopg2.connect(DATABASE_URL)
    cur = conn.cursor()

    for sheet in ["accounts", "orders", "tickets"]:
        load_sheet(cur, wb, sheet)

    conn.commit()
    cur.close()
    conn.close()
    print("Done.")


if __name__ == "__main__":
    main()
